"""Import three bank statements without rewriting workbook extensions.
Keep original source values in extra columns and never guess counterparties.
Run once on the empty transaction template with the three supplied files.
"""
from pathlib import Path
from zipfile import ZipFile
import re,os,tempfile
from xml.sax.saxutils import escape
import openpyxl
from openpyxl.utils import get_column_letter
B=Path('Mazandasht_Tamplate V00')
def number(v):
 if v is None or v=='':return None
 return int(str(v).strip().replace(',',''))
def txt(v):
 if v is None:return ''
 return str(v)
def cell(ref,value,style=None):
 st=f' s="{style}"' if style else ''
 if value is None or value=='':return f'<c r="{ref}"{st}/>'
 if isinstance(value,(int,float)):return f'<c r="{ref}"{st}><v>{value}</v></c>'
 return f'<c r="{ref}"{st} t="inlineStr"><is><t xml:space="preserve">{escape(str(value))}</t></is></c>'
def replace(row,ref,value):
 pattern=rf'<c\s+[^>]*\br="{ref}"[^>]*?(?:/>|>.*?</c>)'
 def sub(m):
  st=re.search(r'\bs="(\d+)"',m.group());return cell(ref,value,st.group(1) if st else None)
 row,n=re.subn(pattern,sub,row,count=1,flags=re.S);assert n==1,ref
 return row
def rewrite(path, transforms):
 with ZipFile(path) as z:
  fd,tmp=tempfile.mkstemp(dir=path.parent,suffix='.xlsx');os.close(fd)
  try:
   with ZipFile(tmp,'w') as out:
    for i in z.infolist():
     data=z.read(i.filename)
     if i.filename in transforms:data=transforms[i.filename](data.decode()).encode()
     out.writestr(i,data)
   os.replace(tmp,path)
  finally:
   if os.path.exists(tmp):os.unlink(tmp)
def rowmap(xml,func):return re.sub(r'<row\b[^>]*\br="(\d+)"[^>]*(?:/>|>.*?</row>)',lambda m:func(int(m.group(1)),m.group()),xml,flags=re.S)
records=[]
for file,bank,start in [('صادرات 9009.xlsx','BANK01 – صادرات 9009',4),('حساب کشاورزی.xlsx','BANK04 – کشاورزی',2),('بانک مهر ایران.xlsx','BANK06 – مهر ایران',17)]:
 sh=openpyxl.load_workbook(file,read_only=True,data_only=True).active
 for rowno,raw in enumerate(sh.iter_rows(values_only=True),1):
  if rowno<start:continue
  if bank.startswith('BANK06'):
   parts=txt(raw[0]).split(' ',1)
   if len(parts)!=2 or not parts[0].startswith('1405/'):continue
   date,t=parts; kind=txt(raw[1]); amount=number(raw[2]);dep=amount if kind=='واریز' else None;wit=amount if kind=='برداشت' else None;bal=number(raw[3]);note=txt(raw[4]);reverse=False
   if kind not in ('واریز','برداشت'):raise ValueError((file,rowno,kind))
  else:
   date=txt(raw[1]);t=txt(raw[2]);reverse=True
   if not date.startswith('1405/'):continue
   if bank.startswith('BANK01'):dep=number(raw[7]);wit=number(raw[8]);bal=number(raw[9]);note=txt(raw[4])
   else:dep=number(raw[4]);wit=number(raw[5]);bal=number(raw[6]);note=txt(raw[3])
  dep=dep or None;wit=wit or None
  assert (dep is None)!=(wit is None),(file,rowno,dep,wit)
  assert re.fullmatch(r'\d{4}/\d{2}/\d{2}',date) and re.match(r'\d{2}:\d{2}',t),(file,rowno,date,t)
  records.append(dict(file=file,bank=bank,date=date,time=t,dep=dep,wit=wit,bal=bal,note=note,rowno=rowno,raw=raw,reverse=reverse))
assert len(records)==378
records.sort(key=lambda a:(a['date'],a['time'],a['bank'], -a['rowno'] if a['reverse'] else a['rowno']))
opening={}
for bank in ('BANK01','BANK04','BANK06'):
 a=next(r for r in records if r['bank'].startswith(bank));opening[bank]=a['bal']-(a['dep'] or 0)+(a['wit'] or 0)
print('opening estimates',opening)

# Update just the three opening balances in INFO's Banks sheet.
info=B/'INFO.xlsx'
w=openpyxl.load_workbook(info,read_only=True);banks_index=w.sheetnames.index('Banks')+1;w.close()
def infoxml(xml):
 def process(r,row):
  if r in (7,10,12):row=replace(row,f'C{r}',opening[{7:'BANK01',10:'BANK04',12:'BANK06'}[r]])
  return row
 return rowmap(xml,process)
rewrite(info,{f'xl/worksheets/sheet{banks_index}.xml':infoxml})

txn=B/'Transactions.xlsx'
w=openpyxl.load_workbook(txn,read_only=True)
entry_index=w.sheetnames.index('Entry')+1
assert not any(c.value is not None and c.data_type!='f' for row in w['Entry'].iter_rows(min_row=9) for c in row)
w.close()
def entry(xml):
 def process(r,row):
  if r==7:
   extras=cell('T7','فایل منبع',1)+cell('U7','ردیف منبع',1)+''.join(cell(f'{get_column_letter(22+i)}7',f'منبع {get_column_letter(i+1)}',1) for i in range(12))
   return row.replace('</row>',extras+'</row>')
  if r==8:
   extras=cell('T8','Source file',2)+cell('U8','Source row',2)+''.join(cell(f'{get_column_letter(22+i)}8',f'Source {get_column_letter(i+1)}',2) for i in range(12))
   return row.replace('</row>',extras+'</row>')
  if 9<=r<9+len(records):
   x=records[r-9];inputs={'C':x['date'],'D':x['time'][:5],'E':x['bank'],'J':'واریز' if x['dep'] is not None else 'برداشت','M':x['dep'],'N':x['wit'],'O':x['bal'],'R':x['note']}
   for col,value in inputs.items():row=replace(row,f'{col}{r}',value)
   extras=cell(f'T{r}',x['file'])+cell(f'U{r}',x['rowno'])+''.join(cell(f'{get_column_letter(22+i)}{r}',txt(x['raw'][i]) if i<len(x['raw']) else None) for i in range(12))
   return row.replace('</row>',extras+'</row>')
  return row
 xml=rowmap(xml,process)
 xml=xml.replace('ref="A1:S2008"','ref="A1:AG2008"',1)
 xml=xml.replace('</cols>','<col min="20" max="21" width="22" customWidth="1"/><col min="22" max="33" width="24" customWidth="1"/></cols>',1)
 # Fix the broken date validator's missing OR and subgroup list's accidental Entry-sheet references.
 xml=xml.replace('AND(LEN($C9)=10,MID($C9,5,1)="/",MID($C9,8,1)="/",ISNUMBER(VALUE(LEFT($C9,4))),ISNUMBER(VALUE(MID($C9,6,2))),ISNUMBER(VALUE(MID($C9,9,2)))),AND(ISNUMBER($C9),LEN($C9)=8)', 'OR(AND(LEN($C9)=10,MID($C9,5,1)="/",MID($C9,8,1)="/",ISNUMBER(VALUE(LEFT($C9,4))),ISNUMBER(VALUE(MID($C9,6,2))),ISNUMBER(VALUE(MID($C9,9,2)))),AND(ISNUMBER($C9),LEN($C9)=8))')
 old='IFERROR(OFFSET($BG$6,0,MATCH($K9,$BG$5:$BZ$5,0)-1,MAX(COUNTIF(OFFSET($BG$6,0,MATCH($K9,$BG$5:$BZ$5,0)-1,100,1),"?*"),1),1),OFFSET($BG$6,0,0,1,1))'
 new='IFERROR(OFFSET(INDIRECT("Sync!$BG$6"),0,MATCH($K9,INDIRECT("Sync!$BG$5:$BZ$5"),0)-1,MAX(COUNTIF(OFFSET(INDIRECT("Sync!$BG$6"),0,MATCH($K9,INDIRECT("Sync!$BG$5:$BZ$5"),0)-1,100,1),"?*"),1),1),OFFSET(INDIRECT("Sync!$BG$6"),0,0,1,1))'
 assert old in xml;xml=xml.replace(old,new)
 xml=xml.replace('AND(ISNUMBER($C9),LEN($C9)=8)))','AND(ISNUMBER($C9),LEN($C9)=8))',1)
 return xml
def workbook(xml):
 xml=re.sub(r'<calcPr\b[^>]*/>','<calcPr calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>',xml)
 # cheque details and raw bank columns remain on sheet but outside print area.
 xml=xml.replace("'Entry'!$A$1:$S$2008","'Entry'!$A$1:$R$2008")
 xml=xml.replace('Entry!$A$1:$S$2008','Entry!$A$1:$R$2008').replace("&apos;Entry&apos;!$A$1:$S$2008","&apos;Entry&apos;!$A$1:$R$2008")
 return xml
rewrite(txn,{f'xl/worksheets/sheet{entry_index}.xml':entry,'xl/workbook.xml':workbook})
print('imported',len(records))
